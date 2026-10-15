"""
Etapa 5 - Relatorio em Excel com graficos (Bolsa_Insights)

Le precos.csv, metricas.csv e insights.txt e gera:
- grafico_desempenho.png    : precos em base 100 (todos comecam em 100)
- grafico_retorno_risco.png : dispersao volatilidade x retorno 12m
- relatorio_bolsa.xlsx      : abas Resumo, Graficos, Metricas e Precos
"""
import matplotlib

matplotlib.use("Agg")  # gera imagens sem abrir janela
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Font

# ---------------- Configuracao ----------------
ARQUIVO_PRECOS = "precos.csv"
ARQUIVO_METRICAS = "metricas.csv"
ARQUIVO_INSIGHTS = "insights.txt"
ARQUIVO_EXCEL = "relatorio_bolsa.xlsx"
GRAFICO_DESEMPENHO = "grafico_desempenho.png"
GRAFICO_RETORNO_RISCO = "grafico_retorno_risco.png"
REFERENCIA = "IBOV"

COLUNAS_PCT = [
    "Retorno_1m", "Retorno_3m", "Retorno_12m", "Dist_MM50",
    "Volatilidade_Anual", "Dist_Maxima", "Drawdown_Max", "Pct_Dias_Alta",
]


# ---------------- Graficos ----------------
def grafico_desempenho(precos, caminho):
    """Linha com todos os ativos em base 100 no primeiro pregao."""
    base100 = precos / precos.iloc[0] * 100

    fig, ax = plt.subplots(figsize=(10, 5))
    for ativo in base100.columns:
        eh_ref = ativo == REFERENCIA
        ax.plot(
            base100.index, base100[ativo], label=ativo,
            linewidth=2.5 if eh_ref else 1.5,
            linestyle="--" if eh_ref else "-",
            color="black" if eh_ref else None,
        )
    ax.axhline(100, color="gray", linewidth=0.8)
    ax.set_title("Desempenho em 12 meses (base 100 = primeiro pregao)")
    ax.set_ylabel("Indice (base 100)")
    ax.legend()
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(caminho, dpi=100)
    plt.close(fig)


def grafico_retorno_risco(metricas, caminho):
    """Dispersao: volatilidade anual (X) x retorno 12m (Y)."""
    x = metricas["Volatilidade_Anual"] * 100
    y = metricas["Retorno_12m"] * 100

    fig, ax = plt.subplots(figsize=(7, 5))
    ax.scatter(x, y, s=90)
    for ativo in metricas.index:
        ax.annotate(ativo, (x[ativo], y[ativo]), xytext=(7, 5), textcoords="offset points")
    ax.set_title("Retorno x risco (12 meses)")
    ax.set_xlabel("Volatilidade anual (%)")
    ax.set_ylabel("Retorno em 12 meses (%)")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(caminho, dpi=100)
    plt.close(fig)


# ---------------- Excel ----------------
def formatar_metricas(ws):
    """Aplica % nas colunas de taxa e 2 casas nas demais."""
    for celula in ws[1]:
        if celula.column == 1:
            continue
        formato = "0.0%" if celula.value in COLUNAS_PCT else "#,##0.00"
        for linha in range(2, ws.max_row + 1):
            ws.cell(row=linha, column=celula.column).number_format = formato
        ws.column_dimensions[celula.column_letter].width = 18


def formatar_precos(ws):
    ws.column_dimensions["A"].width = 12
    for linha in range(2, ws.max_row + 1):
        ws.cell(row=linha, column=1).number_format = "dd/mm/yyyy"
        for col in range(2, ws.max_column + 1):
            ws.cell(row=linha, column=col).number_format = "#,##0.00"


def salvar_excel(precos, metricas, linhas_insights):
    with pd.ExcelWriter(ARQUIVO_EXCEL, engine="openpyxl") as writer:
        pd.DataFrame({"Insights": linhas_insights}).to_excel(writer, sheet_name="Resumo", index=False)
        pd.DataFrame().to_excel(writer, sheet_name="Graficos")
        metricas.to_excel(writer, sheet_name="Metricas")
        precos.to_excel(writer, sheet_name="Precos")

        ws_resumo = writer.sheets["Resumo"]
        ws_resumo.column_dimensions["A"].width = 130
        for linha in range(2, ws_resumo.max_row + 1):
            celula = ws_resumo.cell(row=linha, column=1)
            if celula.value and str(celula.value).startswith(("[", "===")):
                celula.font = Font(bold=True)

        ws_graf = writer.sheets["Graficos"]
        ws_graf.add_image(Image(GRAFICO_DESEMPENHO), "A1")
        ws_graf.add_image(Image(GRAFICO_RETORNO_RISCO), "A28")

        formatar_metricas(writer.sheets["Metricas"])
        formatar_precos(writer.sheets["Precos"])


def main():
    precos = pd.read_csv(ARQUIVO_PRECOS, index_col="Data", parse_dates=True)
    metricas = pd.read_csv(ARQUIVO_METRICAS, index_col="Ativo")
    with open(ARQUIVO_INSIGHTS, encoding="utf-8") as f:
        linhas_insights = f.read().splitlines()

    grafico_desempenho(precos, GRAFICO_DESEMPENHO)
    grafico_retorno_risco(metricas, GRAFICO_RETORNO_RISCO)
    print(f"Graficos salvos: {GRAFICO_DESEMPENHO}, {GRAFICO_RETORNO_RISCO}")

    try:
        salvar_excel(precos, metricas, linhas_insights)
    except PermissionError:
        raise SystemExit(f"ERRO: feche o {ARQUIVO_EXCEL} no Excel e rode de novo.")

    print(f"Relatorio salvo: {ARQUIVO_EXCEL} (abas Resumo, Graficos, Metricas, Precos)")


if __name__ == "__main__":
    main()